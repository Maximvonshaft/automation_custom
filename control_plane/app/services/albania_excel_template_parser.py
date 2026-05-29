from __future__ import annotations

from dataclasses import asdict, dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

DEFAULT_TEMPLATE_SHEET_NAME = "Combine合并订单"
HEADER_ROW_PRIMARY = 1
HEADER_ROW_SECONDARY = 2
FIRST_DATA_ROW = 3


@dataclass(frozen=True)
class TemplateDefaults:
    customs_office: str = "AL111000"
    authorization_reference: str = ""
    means_transport_code_1: str = "50"
    means_transport_country_from: str = "RS"
    truck_registration_plate_number: str = ""
    means_transport_country_to: str = "AL"
    customs_tariff_2: str = "000"
    document_date: str = ""


@dataclass(frozen=True)
class ParsedTemplateGroup:
    group_id: str
    group_index: int
    start_row: int
    end_row: int
    item_count: int
    recipient_name: str
    address: str
    document_reference: str
    document_references: tuple[str, ...]
    customs_tariff_1: str
    statistical_quantity: str
    invoice_value: str
    total_weight: str
    goods_description: str
    declaration: dict[str, str]
    warnings: tuple[str, ...] = ()

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data["document_references"] = list(self.document_references)
        return data


@dataclass(frozen=True)
class TemplateParseResult:
    workbook_path: str
    sheet_name: str
    group_count: int
    groups: tuple[ParsedTemplateGroup, ...]
    warnings: tuple[str, ...] = ()

    def to_dict(self) -> dict[str, Any]:
        return {
            "workbook_path": self.workbook_path,
            "sheet_name": self.sheet_name,
            "group_count": self.group_count,
            "groups": [group.to_dict() for group in self.groups],
            "warnings": list(self.warnings),
        }


HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "row_no": ("Nr.r.",),
    "product": ("Malli/  Produkti", "Malli/ Produkti"),
    "document_reference": ("Kodi",),
    "quantity": ("Sasia",),
    "weight": ("Pesha",),
    "price_usd": ("Çmimi/$", "Cmimi/$"),
    "origin": ("Origjina prej nga vije", "Origjina prej nga vije "),
    "sender": ("Derguesi/ shitësi", "Derguesi/ shitesi"),
    "recipient": ("Pranuesi",),
    "address": ("Adresa",),
    "remarks": ("Verejtje",),
    "raw_hscode": ("HsCode",),
    "carton_number": ("Carton number",),
    "frequency_count": ("出现次数", "Frequency / Count"),
    "consolidated_value": ("合并货值", "Consolidated Value"),
    "value_eur": ("欧元货值", "Value (EUR)"),
    "value_all": ("Value (ALL)",),
    "duty_rate": ("关税税率", "Duty Rate"),
    "tariff_code": ("税率HS code", "HS Code / Tariff Code"),
    "customs_duty": ("关税", "Customs Duty"),
    "vat": ("增值税", "VAT"),
    "total_tax_usd": ("总税额（美元）", "Total Tax (USD)"),
    "total_tax_amount": ("总税额", "Total Tax Amount"),
    "notes": ("提示", "Remarks / Notes"),
    "total_weight": ("总重量(kg)", "Total Weight (kg)"),
    "description_goods": ("货物描述",),
    "customs_description": ("报关货描",),
}


def _norm_header(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


def _build_column_map(sheet: Any) -> dict[str, int]:
    primary: dict[str, int] = {}
    secondary: dict[str, int] = {}
    for column in range(1, sheet.max_column + 1):
        row1 = _norm_header(sheet.cell(HEADER_ROW_PRIMARY, column).value)
        row2 = _norm_header(sheet.cell(HEADER_ROW_SECONDARY, column).value)
        if row1:
            primary[row1] = column
        if row2:
            secondary[row2] = column

    column_map: dict[str, int] = {}
    for key, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            normalized = _norm_header(alias)
            if normalized in primary:
                column_map[key] = primary[normalized]
                break
        if key in column_map:
            continue
        for alias in aliases:
            normalized = _norm_header(alias)
            if normalized in secondary:
                column_map[key] = secondary[normalized]
                break
    return column_map


def _cell(sheet: Any, row: int, column_map: dict[str, int], key: str) -> Any:
    column = column_map.get(key)
    if column is None:
        return None
    return sheet.cell(row, column).value


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _decimal_text(value: Any) -> str:
    if value is None or str(value).strip() == "":
        return ""
    if isinstance(value, float | int):
        decimal = Decimal(str(value))
    else:
        text = str(value).strip().replace(",", "")
        try:
            decimal = Decimal(text)
        except InvalidOperation:
            return text
    if decimal == decimal.to_integral():
        return str(decimal.quantize(Decimal(1)))
    return format(decimal.normalize(), "f")


def _is_group_header(sheet: Any, row: int, column_map: dict[str, int]) -> bool:
    frequency = _cell(sheet, row, column_map, "frequency_count")
    consolidated_value = _cell(sheet, row, column_map, "consolidated_value")
    tariff_code = _cell(sheet, row, column_map, "tariff_code")
    total_weight = _cell(sheet, row, column_map, "total_weight")
    return all(
        value not in (None, "")
        for value in (frequency, consolidated_value, tariff_code, total_weight)
    )


def _non_empty_document_refs(
    sheet: Any,
    rows: range,
    column_map: dict[str, int],
) -> tuple[str, ...]:
    refs: list[str] = []
    for row in rows:
        value = _text(_cell(sheet, row, column_map, "document_reference"))
        if value:
            refs.append(value)
    return tuple(refs)


def _sum_decimal_column(sheet: Any, rows: range, column_map: dict[str, int], key: str) -> str:
    total = Decimal("0")
    found = False
    for row in rows:
        value = _cell(sheet, row, column_map, key)
        if value in (None, ""):
            continue
        try:
            total += Decimal(str(value).replace(",", ""))
            found = True
        except InvalidOperation:
            continue
    return _decimal_text(total) if found else ""


def _sanitize_group_id(value: str, group_index: int) -> str:
    base = "".join(char if char.isalnum() else "_" for char in value).strip("_")
    return base or f"group_{group_index:03d}"


def _build_group(
    sheet: Any,
    start_row: int,
    end_row: int,
    group_index: int,
    column_map: dict[str, int],
    defaults: TemplateDefaults,
) -> ParsedTemplateGroup:
    rows = range(start_row, end_row + 1)
    document_refs = _non_empty_document_refs(sheet, rows, column_map)
    document_reference = document_refs[0] if document_refs else ""
    group_id = _sanitize_group_id(document_reference, group_index)

    total_quantity = _sum_decimal_column(sheet, rows, column_map, "quantity")
    total_weight = _decimal_text(
        _cell(sheet, start_row, column_map, "total_weight")
    ) or _sum_decimal_column(sheet, rows, column_map, "weight")
    invoice_value = _decimal_text(_cell(sheet, start_row, column_map, "consolidated_value"))
    tariff_code = _text(_cell(sheet, start_row, column_map, "tariff_code"))
    customs_description = _text(
        _cell(sheet, start_row, column_map, "customs_description")
    ) or _text(_cell(sheet, start_row, column_map, "description_goods"))
    origin = _text(_cell(sheet, start_row, column_map, "origin"))
    recipient = _text(_cell(sheet, start_row, column_map, "recipient"))
    address = _text(_cell(sheet, start_row, column_map, "address"))

    warnings: list[str] = []
    if not defaults.authorization_reference:
        warnings.append(
            "authorization_reference default is empty; "
            "compiler will reject until provided."
        )
    if not defaults.truck_registration_plate_number:
        warnings.append(
            "truck_registration_plate_number default is empty; "
            "compiler will reject until provided."
        )
    if not document_reference:
        warnings.append("No document reference / Kodi found for this group.")

    declaration = {
        "zyrat_doganore": defaults.customs_office,
        "authorization_reference": defaults.authorization_reference,
        "means_transport_code_1": defaults.means_transport_code_1,
        "means_transport_country_from": defaults.means_transport_country_from,
        "truck_registration_plate_number": defaults.truck_registration_plate_number,
        "means_transport_country_to": defaults.means_transport_country_to,
        "customs_tariff_1": tariff_code,
        "customs_tariff_2": defaults.customs_tariff_2,
        "statistical_quantity": total_quantity,
        "goods_description": customs_description,
        "origjina": origin,
        "gross_weight": total_weight,
        "net_weight": total_weight,
        "invoice_value": invoice_value,
        "currency": "",
        "document_reference": document_reference,
        "document_date": defaults.document_date,
        "document_reference_all": "|".join(document_refs),
        "recipient_name": recipient,
        "recipient_address": address,
    }

    return ParsedTemplateGroup(
        group_id=group_id,
        group_index=group_index,
        start_row=start_row,
        end_row=end_row,
        item_count=end_row - start_row + 1,
        recipient_name=recipient,
        address=address,
        document_reference=document_reference,
        document_references=document_refs,
        customs_tariff_1=tariff_code,
        statistical_quantity=total_quantity,
        invoice_value=invoice_value,
        total_weight=total_weight,
        goods_description=customs_description,
        declaration=declaration,
        warnings=tuple(warnings),
    )


def parse_broker_excel_template(
    workbook_path: Path,
    *,
    defaults: TemplateDefaults | None = None,
    sheet_name: str = DEFAULT_TEMPLATE_SHEET_NAME,
) -> TemplateParseResult:
    defaults = defaults or TemplateDefaults()
    path = workbook_path.expanduser()
    workbook = load_workbook(path, data_only=True)
    if sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
    else:
        sheet = workbook.active

    column_map = _build_column_map(sheet)
    warnings: list[str] = []
    required_headers = [
        "document_reference",
        "quantity",
        "frequency_count",
        "consolidated_value",
        "tariff_code",
        "total_weight",
        "customs_description",
    ]
    missing_required_headers = sorted(key for key in required_headers if key not in column_map)
    if missing_required_headers:
        raise ValueError(
            "Broker Excel template is missing required headers: "
            f"{missing_required_headers}"
        )

    group_rows = [
        row
        for row in range(FIRST_DATA_ROW, sheet.max_row + 1)
        if _is_group_header(sheet, row, column_map)
    ]
    if not group_rows:
        raise ValueError(
            "Broker Excel template contains no detected consolidated declaration groups."
        )

    groups: list[ParsedTemplateGroup] = []
    for index, start_row in enumerate(group_rows, start=1):
        next_start = group_rows[index] if index < len(group_rows) else sheet.max_row + 1
        end_row = next_start - 1
        groups.append(_build_group(sheet, start_row, end_row, index, column_map, defaults))

    return TemplateParseResult(
        workbook_path=str(path),
        sheet_name=sheet.title,
        group_count=len(groups),
        groups=tuple(groups),
        warnings=tuple(warnings),
    )
