from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SUPPORTED_SHEETS = {"Sheet1", "Combine", "Simplified", "Separate"}
REQUIRED_HEADERS = {
    "awb",
    "hs_code",
    "quantity",
    "gross_weight",
    "invoice_value",
    "origin",
    "description",
}
HEADER_ALIASES = {
    "awb": {
        "awb",
        "air waybill",
        "airway bill",
        "document reference",
        "document_reference",
        "kodi",
    },
    "hs_code": {"hs", "hs code", "hs_code", "commodity code", "tariff code", "tax hs code"},
    "quantity": {"quantity", "qty", "packages", "pieces"},
    "statistical_quantity": {"statistical quantity", "statistical_quantity", "stat qty"},
    "gross_weight": {"gross weight", "gross_weight", "weight", "total weight", "gw"},
    "net_weight": {"net weight", "net_weight", "nw"},
    "invoice_value": {
        "invoice value",
        "invoice_value",
        "value",
        "customs value",
        "consolidated value",
    },
    "origin": {"origin", "country of origin", "origin country"},
    "description": {
        "description",
        "goods description",
        "item description",
        "declaration description",
    },
}


@dataclass(frozen=True)
class ManifestTemplate:
    template_type: str
    rows: list[dict[str, Any]]
    header_row: int | None = None


def classify_workbook(workbook: dict[str, list[dict[str, Any]]]) -> str:
    available = [sheet for sheet in workbook if sheet in SUPPORTED_SHEETS]
    if not available:
        raise ValueError("Unsupported manifest template: no supported sheet found")
    if "Combine" in available:
        return "Combine"
    if "Separate" in available:
        return "Separate"
    if "Simplified" in available:
        return "Simplified"
    return "Sheet1"


def _normalize_header(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").split())


def _canonical_header(value: Any) -> str | None:
    normalized = _normalize_header(value)
    for canonical, aliases in HEADER_ALIASES.items():
        if normalized in aliases:
            return canonical
    return None


def _canonicalize_row(row: Mapping[str, Any]) -> dict[str, Any]:
    canonical: dict[str, Any] = {}
    for key, value in row.items():
        mapped = _canonical_header(key) or str(key).strip().lower()
        canonical[mapped] = value
    return canonical


def validate_headers(row: Mapping[str, Any]) -> list[str]:
    normalized = set(_canonicalize_row(row))
    return sorted(REQUIRED_HEADERS - normalized)


def parse_manifest(
    workbook: dict[str, list[dict[str, Any]]],
) -> tuple[ManifestTemplate, dict[str, Any]]:
    template_type = classify_workbook(workbook)
    rows = [_canonicalize_row(row) for row in workbook.get(template_type, [])]
    errors: list[str] = []
    warnings: list[str] = []
    if not rows:
        errors.append(f"{template_type} sheet is empty")
    else:
        missing = validate_headers(rows[0])
        if missing:
            errors.append(f"Missing required header(s): {', '.join(missing)}")
    accepted = not errors
    report = {
        "template_type": template_type,
        "accepted": accepted,
        "row_count": len(rows),
        "errors": errors,
        "warnings": warnings,
    }
    if not accepted:
        raise ValueError(report["errors"][0])
    return ManifestTemplate(template_type=template_type, rows=rows), report


def _sheet_rows(sheet: Any) -> list[list[Any]]:
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def _detect_header_row(rows: Sequence[Sequence[Any]]) -> tuple[int, list[str]]:
    best_index = -1
    best_headers: list[str] = []
    best_score = 0
    for index, row in enumerate(rows[:20]):
        headers = [_canonical_header(value) or "" for value in row]
        score = len(set(headers) & REQUIRED_HEADERS)
        if score > best_score:
            best_index = index
            best_headers = headers
            best_score = score
    if best_index < 0 or best_score < 3:
        raise ValueError("Could not detect manifest header row")
    return best_index, best_headers


def _rows_from_worksheet(sheet: Any) -> tuple[list[dict[str, Any]], int]:
    raw_rows = _sheet_rows(sheet)
    header_index, headers = _detect_header_row(raw_rows)
    parsed_rows: list[dict[str, Any]] = []
    for raw_row in raw_rows[header_index + 1 :]:
        if all(value is None or str(value).strip() == "" for value in raw_row):
            continue
        parsed: dict[str, Any] = {}
        for position, header in enumerate(headers):
            if header:
                parsed[header] = raw_row[position] if position < len(raw_row) else None
        parsed_rows.append(parsed)
    return parsed_rows, header_index + 1


def load_workbook_rows_from_bytes(content: bytes) -> dict[str, list[dict[str, Any]]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    parsed: dict[str, list[dict[str, Any]]] = {}
    for sheet_name in workbook.sheetnames:
        if sheet_name not in SUPPORTED_SHEETS:
            continue
        rows, _header_row = _rows_from_worksheet(workbook[sheet_name])
        parsed[sheet_name] = rows
    return parsed


def load_workbook_rows_from_path(path: Path) -> dict[str, list[dict[str, Any]]]:
    return load_workbook_rows_from_bytes(path.read_bytes())


def parse_manifest_bytes(content: bytes) -> tuple[ManifestTemplate, dict[str, Any]]:
    return parse_manifest(load_workbook_rows_from_bytes(content))


def parse_manifest_path(path: Path) -> tuple[ManifestTemplate, dict[str, Any]]:
    return parse_manifest_bytes(path.read_bytes())
