from __future__ import annotations

import pytest
from openpyxl import Workbook

from control_plane.app.services.manifest_intake import (
    parse_manifest,
    parse_manifest_bytes,
    parse_manifest_path,
)
from control_plane.app.services.template_classifier import classify_template


def sanitized_workbook(sheet: str = "Combine") -> dict:
    return {
        sheet: [
            {
                "awb": "SANITIZED-AWB-001",
                "hs_code": "610910",
                "quantity": "2",
                "gross_weight": "4.5",
                "net_weight": "4.0",
                "invoice_value": "120.00",
                "origin": "AL",
                "description": "sanitized cotton shirts",
            }
        ]
    }


@pytest.mark.parametrize("sheet", ["Sheet1", "Combine", "Simplified", "Separate"])
def test_classifies_supported_manifest_sheets(sheet):
    contract = classify_template(sanitized_workbook(sheet))
    assert contract["template_type"] == sheet
    assert "awb" in contract["required_headers"]


@pytest.mark.parametrize("sheet", ["Combine", "Separate"])
def test_parses_p0_manifest_templates(sheet):
    manifest, report = parse_manifest(sanitized_workbook(sheet))
    assert manifest.template_type == sheet
    assert report == {
        "template_type": sheet,
        "accepted": True,
        "row_count": 1,
        "errors": [],
        "warnings": [],
    }


def test_rejects_missing_required_headers():
    with pytest.raises(ValueError, match="Missing required header"):
        parse_manifest({"Combine": [{"awb": "SANITIZED-AWB-001"}]})


def test_parses_sanitized_xlsx_fixture_with_header_detection_and_aliases(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Combine"
    sheet.append(["ignored", "metadata"])
    sheet.append(
        [
            "Document Reference",
            "Commodity Code",
            "Qty",
            "Gross Weight",
            "Net Weight",
            "Invoice Value",
            "Country of Origin",
            "Goods Description",
        ]
    )
    sheet.append(
        [
            "SANITIZED-DOC-001",
            "610910",
            2,
            4.5,
            4.0,
            120.0,
            "AL",
            "sanitized cotton shirts",
        ]
    )
    path = tmp_path / "sanitized_manifest.xlsx"
    workbook.save(path)

    manifest, report = parse_manifest_path(path)

    assert report["accepted"] is True
    assert manifest.template_type == "Combine"
    assert manifest.rows[0]["awb"] == "SANITIZED-DOC-001"
    assert manifest.rows[0]["hs_code"] == "610910"
    assert manifest.rows[0]["gross_weight"] == 4.5


def test_parses_sanitized_xlsx_bytes_input(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Separate"
    sheet.append(["AWB", "HS Code", "Quantity", "Gross Weight", "Value", "Origin", "Description"])
    sheet.append(["SANITIZED-DOC-002", "420221", 1, 0.8, 80.0, "AL", "sanitized bag"])
    path = tmp_path / "sanitized_manifest_bytes.xlsx"
    workbook.save(path)

    manifest, report = parse_manifest_bytes(path.read_bytes())

    assert report["accepted"] is True
    assert manifest.template_type == "Separate"
    assert manifest.rows[0]["invoice_value"] == 80.0


def test_albania_frontline_aliases_are_supported(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Combine"
    sheet.append(
        [
            "Kodi",
            "tax HS code",
            "Quantity",
            "total weight",
            "consolidated value",
            "country of origin",
            "declaration description",
        ]
    )
    sheet.append(
        [
            "SANITIZED-KODI-001",
            "610910",
            2,
            4.5,
            120.0,
            "AL",
            "sanitized declaration",
        ]
    )
    path = tmp_path / "sanitized_alias_manifest.xlsx"
    workbook.save(path)

    manifest, report = parse_manifest_path(path)

    assert report["accepted"] is True
    assert manifest.rows[0]["awb"] == "SANITIZED-KODI-001"
    assert manifest.rows[0]["hs_code"] == "610910"
    assert manifest.rows[0]["gross_weight"] == 4.5
    assert manifest.rows[0]["invoice_value"] == 120.0
    assert manifest.rows[0]["origin"] == "AL"
    assert manifest.rows[0]["description"] == "sanitized declaration"
